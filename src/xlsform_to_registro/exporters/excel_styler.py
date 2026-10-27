from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..core.schema import CatalogueSheet, EntitySheet, Variable


class ExcelStyler:
    """Generador y formateador de diseño de registro en Excel usando openpyxl."""

    def __init__(
        self,
        header_color: str = "64b692",
        header_text_color: str = "FFFFFF",
        font_name: str = "Arial",
        header_font_size: int = 11,
        data_font_size: int = 10,
        add_not_applicable_code: bool = True,
    ):
        self.header_color = header_color
        self.header_text_color = header_text_color
        self.font_name = font_name
        self.header_font_size = header_font_size
        self.data_font_size = data_font_size
        self.add_not_applicable_code = add_not_applicable_code

        # OpenPyXL Styles
        self.header_fill = PatternFill(
            start_color=self.header_color,
            end_color=self.header_color,
            fill_type="solid"
        )
        self.header_font = Font(
            name=self.font_name,
            size=self.header_font_size,
            bold=True,
            color=self.header_text_color
        )
        self.section_fill = PatternFill(
            start_color=self.header_color,
            end_color=self.header_color,
            fill_type="solid"
        )
        self.section_font = Font(
            name=self.font_name,
            size=self.header_font_size,
            bold=True,
            color=self.header_text_color
        )
        self.data_font = Font(
            name=self.font_name,
            size=self.data_font_size,
            bold=False,
            color="000000"
        )

        thin_side = Side(style="thin", color="B0B0B0")
        self.data_border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side
        )

        self.align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.align_section = Alignment(horizontal="left", vertical="center", indent=1)
        self.align_var_name = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self.align_var_label = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self.align_value = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.align_value_label = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self.align_universe = Alignment(horizontal="left", vertical="center", wrap_text=True)

        self.col_headers = [
            "Nombre de la variable",
            "Etiqueta de la variable",
            "Valor",
            "Etiqueta de valor",
            "Universo de referencia"
        ]

        self.col_widths = {
            "A": 28,
            "B": 36,
            "C": 12,
            "D": 38,
            "E": 28
        }

    def create_workbook(
        self,
        entity_sheets: List[EntitySheet],
        catalogue_sheets: List[CatalogueSheet]
    ) -> openpyxl.Workbook:
        """Crea el libro Excel y todas sus pestañas formateadas."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # type: ignore # Remover hoja por defecto

        catalogue_names = {cat.list_name: cat.sheet_name for cat in catalogue_sheets}

        # 1. Crear hojas de entidades
        for entity in entity_sheets:
            clean_sheet_name = entity.name[:31]
            ws = wb.create_sheet(clean_sheet_name)
            self._setup_entity_sheet(ws)
            self._populate_entity_sheet(ws, entity, catalogue_names)

        # 2. Crear hojas de catálogos complementarios
        for cat in catalogue_sheets:
            ws = wb.create_sheet(cat.sheet_name[:31])
            self._setup_catalogue_sheet(ws, cat)

        return wb

    def _setup_entity_sheet(self, ws: Any) -> None:
        """Configura cuadrícula, encabezados y anchos de columna en una hoja de entidad."""
        ws.views.sheetView[0].showGridLines = True
        ws.row_dimensions[1].height = 28

        for col_idx, h_text in enumerate(self.col_headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h_text)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_header
            cell.border = self.data_border

        for col_letter, width in self.col_widths.items():
            ws.column_dimensions[col_letter].width = width

    def _populate_entity_sheet(
        self,
        ws: Any,
        entity: EntitySheet,
        catalogue_names: Dict[str, str]
    ) -> None:
        """Rellena secciones, variables, opciones y combina celdas verticalmente."""
        current_row = 2

        for section in entity.sections:
            # Fila de sección
            ws.row_dimensions[current_row].height = 24
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
            for col in range(1, 6):
                c = ws.cell(row=current_row, column=col)
                c.fill = self.section_fill
                c.border = self.data_border

            cell_a = ws.cell(row=current_row, column=1, value=section.title)
            cell_a.font = self.section_font
            cell_a.alignment = self.align_section
            current_row += 1

            for var in section.variables:
                current_row = self._write_variable(ws, var, current_row, catalogue_names)

    def _get_continuous_description(self, var: Variable) -> str:
        """Determina la descripción del valor para variables numéricas o continuas."""
        t = var.qtype.split()[0].lower()
        if t in ["integer", "int"]:
            return "Valor numérico entero"
        elif t in ["decimal"]:
            return "Valor numérico decimal"
        elif t in ["text", "string"]:
            return "Texto alfanumérico libre"
        elif t in ["date", "today"]:
            return "Fecha en formato AAAA-MM-DD"
        elif t in ["time"]:
            return "Hora en formato HH:MM"
        elif t in ["datetime"]:
            return "Fecha y hora en formato ISO"
        elif t in ["geopoint"]:
            return "Coordenadas GPS (Latitud, Longitud, Altitud, Precisión)"
        elif t in ["calculate"]:
            calc_info = f": {var.calculation}" if var.calculation else ""
            return f"Variable calculada internamente{calc_info}"
        return "Respuesta continua / registro abierto"

    def _write_variable(
        self,
        ws: Any,
        var: Variable,
        current_row: int,
        catalogue_names: Dict[str, str]
    ) -> int:
        """Escribe una variable y sus opciones de respuesta, aplicando combinaciones verticales."""
        # 1. Caso: Variable categórica referenciada en hoja de catálogo independiente
        if var.choice_list and var.choice_list in catalogue_names:
            cat_sheet = catalogue_names[var.choice_list]
            values_labels = [
                ("...", f"Código(s) según clasificador (Ver hoja '{cat_sheet}')")
            ]
        # 2. Caso: Variable categórica con opciones locales
        elif var.choices:
            values_labels = [(c.name, c.label) for c in var.choices]
            if var.has_other:
                values_labels.append(("other", "Otro/a (especificar)"))
            # Si tiene filtro (salto de pregunta) y está habilitado el código 0
            if self.add_not_applicable_code and var.relevant:
                values_labels.append((0, "No corresponde / No aplica"))
        # 3. Caso: Variable externa
        elif var.is_external:
            values_labels = [
                ("...", f"Opciones provistas externamente ({var.choice_list or 'archivo CSV'})")
            ]
        # 4. Caso: Variable abierta / continua
        else:
            values_labels = [("...", self._get_continuous_description(var))]

        n_rows = len(values_labels)
        start_row = current_row
        end_row = current_row + n_rows - 1

        for idx, (val, val_lbl) in enumerate(values_labels):
            r = start_row + idx
            ws.row_dimensions[r].height = 20

            c_name = ws.cell(row=r, column=1, value=var.name if idx == 0 else None)
            c_label = ws.cell(row=r, column=2, value=var.label if idx == 0 else None)
            c_val = ws.cell(row=r, column=3, value=val)
            c_val_lbl = ws.cell(row=r, column=4, value=val_lbl)
            c_univ = ws.cell(row=r, column=5, value=var.universe if idx == 0 else None)

            for c, align in [
                (c_name, self.align_var_name),
                (c_label, self.align_var_label),
                (c_val, self.align_value),
                (c_val_lbl, self.align_value_label),
                (c_univ, self.align_universe),
            ]:
                c.font = self.data_font
                c.border = self.data_border
                c.alignment = align

        if n_rows > 1:
            ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
            ws.merge_cells(start_row=start_row, start_column=5, end_row=end_row, end_column=5)

        return end_row + 1

    def _setup_catalogue_sheet(self, ws: Any, cat: CatalogueSheet) -> None:
        """Formatea y rellena una hoja complementaria de catálogo."""
        ws.views.sheetView[0].showGridLines = True
        ws.row_dimensions[1].height = 28

        for col_idx, h_text in enumerate(cat.headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h_text)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_header
            cell.border = self.data_border

        r = 2
        for row_vals in cat.rows:
            ws.row_dimensions[r].height = 20
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=r, column=col_idx, value=val)
                cell.font = self.data_font
                cell.border = self.data_border
                cell.alignment = self.align_value if col_idx == 1 else self.align_value_label
            r += 1

        # Ajustar anchos automáticos
        for col_idx in range(1, len(cat.headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                len(str(ws.cell(row, col_idx).value or ""))
                for row in range(1, min(100, ws.max_row + 1))
            )
            ws.column_dimensions[col_letter].width = max(14, min(50, max_len + 4))
